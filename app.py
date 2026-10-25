import io, re
from datetime import date
from pathlib import Path

import streamlit as st
from PIL import Image
from pptx import Presentation
from pptx.util import Inches
import threading
from streamlit_webrtc import webrtc_streamer

# Menyimpan frame kamera terbaru sementara
_latest_frame = {"image": None}
_frame_lock = threading.Lock()


def camera_frame_callback(frame):
    image = frame.to_ndarray(format="rgb24")
    with _frame_lock:
        _latest_frame["image"] = image
    return frame

APP_DIR = Path(__file__).parent
TEMPLATE_PATH = APP_DIR / "template_ranking.pptx"
CATEGORIES = ["GLOBAL", "BLAZER", "DRESS", "JACKET", "KNITWEAR", "PANTS", "SHIRT", "SHORT", "SKIRT", "VEST", "DENIM PANTS", "SKIRT (PAGE 2)"]
DISPLAY_SLIDE_INDEX = {"GLOBAL": 2, "BLAZER": 4, "DRESS": 6, "JACKET": 8, "KNITWEAR": 10, "PANTS": 12, "SHIRT": 14, "SHORT": 16, "SKIRT": 18, "VEST": 20, "DENIM PANTS": 22, "SKIRT (PAGE 2)": 24}
SLOT_X_BY_RANK = [0.93, 3.24, 5.62, 7.99, 10.32, 0.92, 3.23, 5.61, 7.99, 10.31]
SLOT_Y = [1.35, 1.36, 1.36, 1.36, 1.37, 4.28, 4.28, 4.28, 4.28, 4.29]
PHOTO_W = 1.63
PHOTO_H = 2.17
PHOTO_INSET_X = 0.20
PHOTO_INSET_Y = 0.08

def format_period(start, end):
    """Format the selected report period in the template's uppercase English style."""
    if start.year == end.year and start.month == end.month:
        return f"{start.day} – {end.day} {end.strftime('%b').upper()} {end.year}"
    if start.year == end.year:
        return f"{start.day} {start.strftime('%b').upper()} – {end.day} {end.strftime('%b').upper()} {end.year}"
    return f"{start.day} {start.strftime('%b').upper()} {start.year} – {end.day} {end.strftime('%b').upper()} {end.year}"


st.set_page_config(page_title="Ranking Position Ladies - Festival Citylink", page_icon="📸", layout="wide")
st.title("RANKING POSITION LADIES — FESTIVAL CITYLINK")
st.caption("Input foto display, pilih PICT referensi, lalu buat PowerPoint dari template FESTIVAL CITYLINK.")

if "entries" not in st.session_state:
    st.session_state.entries = []

with st.container(border=True):
    st.subheader("Laporan mingguan")
    st.write("Setiap minggu dibuat sebagai laporan baru. Foto dari minggu sebelumnya tidak otomatis dibawa ke laporan baru.")
    if st.button("Mulai minggu baru / kosongkan semua foto", type="secondary", width="stretch", key="start_new_week"):
        st.session_state.entries = []
        st.success("Daftar foto sudah dikosongkan. Silakan input foto untuk minggu yang baru.")
        st.rerun()

with st.expander("1. Pilih sumber PICT", expanded=True):
    st.write("Pilih foto PICT dari folder komputer atau Google Drive yang sudah disinkronkan. PICT hanya menjadi referensi pilihan.")
    pict_files = st.file_uploader("Upload beberapa gambar PICT sekaligus", type=["jpg", "jpeg", "png", "webp"], accept_multiple_files=True, key="pict_upload")
    pict_names = [f.name for f in pict_files] if pict_files else []
    if pict_names:
        st.success(f"{len(pict_names)} gambar PICT tersedia untuk dipilih.")
    else:
        st.info("Belum ada PICT yang dipilih. Anda tetap bisa memasukkan foto display.")

st.divider()
st.subheader("2. Input foto display")
left, right = st.columns([1, 1])
with left:
    d1, d2 = st.columns(2)
    with d1:
        start_date = st.date_input("Tanggal mulai laporan", value=date(2026, 10, 19), format="DD/MM/YYYY")
    with d2:
        end_date = st.date_input("Tanggal akhir laporan", value=date(2026, 10, 25), format="DD/MM/YYYY")
    if end_date < start_date:
        st.error("Tanggal akhir tidak boleh lebih awal dari tanggal mulai.")
    category = st.selectbox("Kategori", CATEGORIES, index=0)
    rank = st.selectbox("Ranking Position", list(range(1, 11)), format_func=lambda x: f"RK {x}")
    article = st.text_input("Kode/nama artikel (opsional)", placeholder="Masukkan artikel jika diperlukan")
    selected_pict = st.selectbox("Pilih PICT referensi", pict_names) if pict_names else None
with right:
   
    st.write("**Foto display asli**")

    webrtc_streamer(
        key="display_camera",
        video_frame_callback=camera_frame_callback,
        media_stream_constraints={
            "video": {"facingMode": "environment"},
            "audio": False,
        },
        async_processing=True,
    )

    if st.button("Ambil foto dari kamera belakang", key="capture_display_photo"):
        with _frame_lock:
            latest = _latest_frame["image"]
            captured = latest.copy() if latest is not None else None

        if captured is not None:
            buffer = io.BytesIO()
            Image.fromarray(captured).save(buffer, format="JPEG", quality=90)
            st.session_state["display_camera_photo"] = buffer.getvalue()
            st.success("Foto berhasil diambil.")
        else:
            st.warning("Kamera belum siap. Tunggu sampai gambar muncul.")

    uploaded_photo = st.file_uploader(
        "Atau pilih foto display dari perangkat",
        type=["jpg", "jpeg", "png", "webp"],
        key="display_upload",
    )

    if uploaded_photo:
        display_photo = uploaded_photo
    else:
        saved_photo = st.session_state.get("display_camera_photo")
        display_photo = io.BytesIO(saved_photo) if saved_photo else None

    if display_photo:
        st.image(
            display_photo,
            caption="Foto display asli yang akan masuk ke PowerPoint",
            width="stretch",
        )


if st.button("Simpan foto ke daftar ranking", type="primary", width="stretch"):
    if end_date < start_date:
        st.error("Perbaiki periode tanggal terlebih dahulu.")
    elif not display_photo:
        st.error("Ambil foto display atau pilih file foto terlebih dahulu.")
    else:
        existing = next((e for e in st.session_state.entries if e["category"] == category and e["rank"] == rank), None)
        if existing:
            st.error(f"Slot {category} — RK {rank} sudah terisi. Foto baru tidak disimpan agar tidak menimpa foto sebelumnya. Hapus slot tersebut pada daftar di bawah jika memang ingin menggantinya.")
        else:
            photo_bytes = display_photo.getvalue()
            pict_bytes = next((f.getvalue() for f in pict_files if f.name == selected_pict), None) if selected_pict and pict_files else None
            date_label = format_period(start_date, end_date)
            st.session_state.entries.append({"date": date_label, "category": category, "rank": rank,
                "article": article.strip(), "pict_name": selected_pict or "", "pict_bytes": pict_bytes,
                "display_name": getattr(display_photo, "name", f"camera_{category}_{rank}.jpg"), "display_bytes": photo_bytes})
            st.success(f"Tersimpan: {category} — RK {rank} ({date_label}).")

st.divider()
st.subheader("3. Daftar foto yang akan dimasukkan")
if not st.session_state.entries:
    st.info("Belum ada foto yang disimpan.")
else:
    for e in sorted(st.session_state.entries, key=lambda x: (CATEGORIES.index(x['category']), x['rank'])):
        c1, c2, c3, c4 = st.columns([1, 3, 3, 1])
        with c1:
            st.image(e["display_bytes"], width="stretch")
        with c2:
            st.write(f"**{e['category']} — RK {e['rank']}**")
            st.caption(f"Periode: {e['date']} · Artikel: {e['article'] or '—'}")
        with c3:
            st.caption(f"PICT referensi: {e['pict_name'] or 'Belum dipilih'}")
        with c4:
            if st.button("Hapus", key=f"del_{e['category']}_{e['rank']}"):
                st.session_state.entries = [x for x in st.session_state.entries if not (x['category'] == e['category'] and x['rank'] == e['rank'])]
                st.rerun()


def make_photo_bytes(raw):
    im = Image.open(io.BytesIO(raw)).convert("RGB")
    target_ratio = PHOTO_W / PHOTO_H
    w, h = im.size
    if w / h > target_ratio:
        new_w = int(h * target_ratio)
        left = (w - new_w) // 2
        im = im.crop((left, 0, left + new_w, h))
    else:
        new_h = int(w / target_ratio)
        top = (h - new_h) // 2
        im = im.crop((0, top, w, top + new_h))
    out = io.BytesIO()
    im.save(out, format="JPEG", quality=92)
    out.seek(0)
    return out


def clear_status_textboxes(slide, rank):
    slot = rank - 1
    x0, y0 = SLOT_X_BY_RANK[slot], SLOT_Y[slot]
    for shape in slide.shapes:
        if not getattr(shape, "has_text_frame", False):
            continue
        if shape.text.strip() not in {"√", "No stock", "NO STOCK", "LP", "LAST PIECES"}:
            continue
        sx, sy = shape.left / 914400, shape.top / 914400
        if abs(sx - (x0 + PHOTO_INSET_X)) < 0.35 and abs(sy - (y0 + PHOTO_INSET_Y)) < 0.35:
            shape.text_frame.clear()


def insert_photo_under_overlays(slide, photo_stream, rank):
    slot = rank - 1
    x0, y0 = SLOT_X_BY_RANK[slot], SLOT_Y[slot]
    pic = slide.shapes.add_picture(photo_stream, Inches(x0 + PHOTO_INSET_X), Inches(y0 + PHOTO_INSET_Y), Inches(PHOTO_W), Inches(PHOTO_H))
    sp_tree = slide.shapes._spTree
    sp_tree.remove(pic._element)
    insert_at = len(sp_tree)
    for i, child in enumerate(sp_tree):
        tag = child.tag.split('}')[-1]
        if tag in ("sp", "graphicFrame", "cxnSp"):
            nv = child.find('.//{http://schemas.openxmlformats.org/presentationml/2006/main}cNvPr')
            name = nv.get('name', '') if nv is not None else ''
            if name.startswith("object"):
                try:
                    off = child.find('.//{http://schemas.openxmlformats.org/drawingml/2006/main}off')
                    x, y = int(off.get('x')) / 914400, int(off.get('y')) / 914400
                    if abs(x - x0) < 0.18 and abs(y - y0) < 0.18 and tag == 'sp':
                        insert_at = i
                        break
                except Exception:
                    pass
    sp_tree.insert(insert_at, pic._element)


def generate_pptx(entries):
    prs = Presentation(str(TEMPLATE_PATH))
    date_text = entries[0]["date"] if entries else format_period(date(2026, 10, 19), date(2026, 10, 25))
    for slide in prs.slides:
        for sh in slide.shapes:
            if getattr(sh, "has_text_frame", False) and re.search(r"\bOCT\s+2026\b", sh.text, flags=re.I):
                # Keep the existing font/layout; replace only the date text in the existing run.
                p = sh.text_frame.paragraphs[0]
                if p.runs:
                    p.runs[0].text = date_text
                    for run in p.runs[1:]:
                        run.text = ""
                else:
                    sh.text_frame.text = date_text
    for entry in entries:
        slide = prs.slides[DISPLAY_SLIDE_INDEX[entry["category"]]]
        clear_status_textboxes(slide, entry["rank"])
        insert_photo_under_overlays(slide, make_photo_bytes(entry["display_bytes"]), entry["rank"])
    out = io.BytesIO()
    prs.save(out)
    out.seek(0)
    return out

if st.session_state.entries:
    st.divider()
    if st.button("Generate PowerPoint (.pptx)", type="primary", width="stretch"):
        try:
            result = generate_pptx(st.session_state.entries)
            st.download_button("Download PowerPoint", data=result.getvalue(), file_name=f"RANKING_POSITION_LADIES_FESTIVAL_CITYLINK_{date.today():%Y%m%d}.pptx", mime="application/vnd.openxmlformats-officedocument.presentationml.presentation", width="stretch")
            st.success("PowerPoint berhasil dibuat. Buka hasilnya untuk memeriksa periode dan penempatan foto.")
        except Exception as ex:
            st.error(f"Belum berhasil membuat PowerPoint: {ex}")
            st.caption("Kirim teks error ini agar bisa diperbaiki dengan tepat.")

with st.expander("Catatan"):
    st.markdown("""
- Template yang digunakan adalah file FESTIVAL CITYLINK asli yang sudah disimpan ulang untuk menormalkan file PowerPoint.
- Periode laporan default: **19–25 OCT 2026** dan bisa diubah melalui tanggal mulai/akhir.
- Gunakan **Mulai minggu baru / kosongkan semua foto** sebelum mengisi laporan minggu berikutnya.
- Jika slot kategori dan ranking sudah terisi, aplikasi menolak foto baru agar tidak menimpa foto yang sudah disimpan; hapus slot secara manual jika ingin menggantinya.
- Foto yang dimasukkan ke slide adalah foto display asli; PICT hanya referensi.
- PICT diunggah dari komputer atau folder Google Drive yang tersinkron. Koneksi Google Drive langsung belum disiapkan.
- Template berisi 26 slide; kategori termasuk DENIM PANTS dan halaman SKIRT kedua.
""")
